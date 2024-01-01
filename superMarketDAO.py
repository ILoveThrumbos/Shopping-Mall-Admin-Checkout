import sqlite3
from openpyxl.chart import BarChart, Reference
from openpyxl import Workbook


class SuperMarketDao:
    def __init__(self, product_db):
        self.cursor = None
        self.db_name = product_db
        self.db = None
        self.connect()
        try:
            cursor = self.db.cursor()
            cursor.execute('''CREATE TABLE IF NOT EXISTS product (
                                       barcode CHAR(3) NOT NULL,
                                       name CHAR(25) NOT NULL,
                                       desc CHAR(50) NOT NULL,
                                       price FLOAT NOT NULL
                                   )''')
            cursor.execute('''CREATE TABLE IF NOT EXISTS transactions (
                                       date CHAR(10) NOT NULL,
                                       barcode CHAR(3) NOT NULL,
                                       price FLOAT NOT NULL,
                                       FOREIGN KEY (barcode) REFERENCES product(barcode)
                                   )''')
            self.db.commit()
        except Exception as e:
            print("Unable to create tables. {}".format(e))
            self.db.rollback()
            raise e
        finally:
            self.close()

    def connect(self):
        if self.db is None:
            self.db = sqlite3.connect(self.db_name)

    def close(self):
        if self.db is not None:
            self.db.close()
            self.db = None

    def addProduct(self):
        try:
            self.connect()
            cursor = self.db.cursor()
            query = '''INSERT INTO product (barcode, name, desc, price) VALUES (?,?,?,?)'''

            # Barcode input validation
            while True:
                barcode = input('Barcode:')
                if not barcode:
                    print("Barcode cannot be blank.")
                elif not barcode.isdigit() or len(barcode) != 3:
                    print("Invalid barcode. Barcode must be a positive integer and exactly 3 characters in length.")
                else:
                    cursor.execute('''SELECT COUNT(*) FROM product WHERE barcode = ?''', (barcode,))
                    count = cursor.fetchone()[0]
                    if count > 0:
                        print('Barcode already exists in the database.')
                    else:
                        break

            # Name input validation
            while True:
                name = input('Name:')
                if not name:
                    print("Name cannot be blank.")
                elif len(name) > 25:
                    print("Name cannot be more than 25 characters long.")
                else:
                    break

            # Description input validation
            while True:
                desc = input('Description:')
                if not desc:
                    print("Description cannot be blank.")
                elif len(desc) > 50:
                    print("Description cannot be more than 50 characters long.")
                else:
                    break

            # Price input validation
            while True:
                price = input('Price:')
                if not price:
                    print("Price cannot be blank.")
                elif price:
                    try:
                        price = float(price)
                        if price <= 0:
                            print("Price must be a positive float value and not 0.")
                        else:
                            break
                    except ValueError:
                        print("Invalid input. Please enter valid numeric numbers.")

            cursor.execute(query, (barcode, name, desc, price))
            self.db.commit()
            cursor.execute('''SELECT * FROM product ORDER BY ROWID DESC LIMIT 1''')
            new_product = cursor.fetchone()
            if new_product:
                print('Data entered successfully.')
                return new_product
            else:
                return None
        except Exception as e:
            print("Unable to insert data into the table. {}".format(e))
            self.db.rollback()
            raise e
        finally:
            self.close()

    def listAllProducts(self):
        try:
            self.connect()
            print('\n---------------------------------\n')
            cursor = self.db.cursor()
            self.db.commit()
            cursor.execute('''SELECT * FROM product ORDER BY barcode ASC''')
            all_rows = cursor.fetchall()
            for row in all_rows:
                print('{0} : {1}, {2}, {3}'.format(row[0], row[1], row[2], row[3]))
            return 'Success! Data displayed successfully.'
        except Exception as e:
            print("Unable to display data. {}".format(e))
            self.db.rollback()
            raise e
        finally:
            self.close()

    def findProduct(self, product):
        try:
            self.connect()
            values = {'barcode': product.get_barcode()}
            cursor = self.db.cursor()
            cursor.execute('''SELECT * FROM product WHERE barcode=:barcode''', values)
            barcode = str(product.get_barcode())
            if not barcode.isdigit():
                raise ValueError("Barcode must be a positive integer.")
            self.db.commit()
            barcode_found = cursor.fetchone()
            if barcode_found:
                print('Barcode successfully found.')
                return barcode_found
            else:
                print('Barcode does not exist in database.')
        except Exception as e:
            print("No barcode is found in database. {}".format(e))
            self.db.rollback()
        finally:
            self.close()

    def listAllTransactions(self):
        try:
            self.connect()
            cursor = self.db.cursor()
            self.db.commit()
            cursor.execute('''SELECT barcode, date, price FROM transactions ORDER BY date ASC''')
            all_rows = cursor.fetchall()
            for row in all_rows:
                print('{0} : {1}, {2}'.format(row[0], row[1], row[2]))
            return 'Success! Data displayed successfully.'

        except Exception as e:
            print("Unable to open transactions table. {}".format(e))
            self.db.rollback()
            raise e
        finally:
            self.close()

    def displayBarchartOfProductsSold(self):
        try:
            self.connect()
            cursor = self.db.cursor()

            # Retrieve product counts from transactions table
            query = '''
            SELECT barcode, COUNT(barcode) as quantity
            FROM transactions
            GROUP BY barcode
            '''
            cursor.execute(query)
            rows = cursor.fetchall()
            product_count = {}
            for row in rows:
                product_count[row[0]] = row[1]

            product_names = []
            quantities = []

            # Retrieve product names from products table using barcode
            for barcode, count in product_count.items():
                query = '''
                SELECT name
                FROM product
                WHERE barcode = '{}'
                '''.format(barcode)
                cursor.execute(query)
                result = cursor.fetchone()
                if result:
                    product_names.append(result[0])
                    quantities.append(count)

            # Create bar chart
            wb = Workbook()
            sheet = wb.active
            sheet.append(['Product', 'Quantity'])
            for i in range(len(product_names)):
                sheet.append([product_names[i], quantities[i]])

            chart = BarChart()
            data = Reference(worksheet=sheet,
                             min_row=2,
                             max_row=len(product_names) + 1,
                             min_col=2,
                             max_col=2)
            chart.add_data(data, titles_from_data=False)

            categories = Reference(worksheet=sheet,
                                   min_row=2,
                                   max_row=len(product_names) + 1,
                                   min_col=1,
                                   max_col=1)
            chart.set_categories(categories)

            chart.x_axis.title = "Product"
            chart.y_axis.title = "Quantity"
            chart.title = "Display a Bar chart of Products sold by quantity"

            sheet.add_chart(chart, "E2")
            wb.save("products_sold.xlsx")
            return 'Success! products_sold document saved.'

        except Exception as e:
            print("Unable to create or save the barchart. {}".format(e))
            raise e
        finally:
            self.close()

    def displayExcelReportOfTransactions(self):
        try:
            self.connect()
            cursor = self.db.cursor()

            query = '''SELECT * FROM transactions ORDER BY date ASC'''
            cursor.execute(query)
            rows = cursor.fetchall()
            wb = Workbook()
            sheet = wb.active
            for row in rows:
                sheet.append(row)
            wb.save("transactions.xls")
            print('transactions report created.')

        except Exception as e:
            print("Unable to open transactions table. {}".format(e))

            self.db.rollback()
            raise e
        finally:
            self.close()
            return rows

    def addTransactionToDB(self, transaction):
        try:
            self.connect()
            cursor = self.db.cursor()
            query = '''
                INSERT INTO transactions (date, barcode, price)
                VALUES (?, ?, ?)
            '''
            cursor.execute(query, (transaction.date, transaction.barcode, transaction.price))
            self.db.commit()
            print('Transaction added to the database.')
        except Exception as e:
            print("Unable to add transaction to the database. {}".format(e))
            self.db.rollback()
            raise e
        finally:
            self.close()





